import os
import re
import PyPDF2
from docx import Document
from flask import Flask, request, render_template, send_file
from openpyxl import Workbook

app = Flask(__name__)

# Function to extract email addresses from text
def extract_emails(text):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    return re.findall(email_pattern, text)

# Function to extract phone numbers from text
def extract_phone_numbers(text):
    phone_pattern = r'\b(?:\+\d{1,2}\s?)?(?:\(\d{3,}\)|\d{3,})[-.\s]?\d{3,}[-.\s]?\d{4}\b'
    return re.findall(phone_pattern, text)

# Function to extract text from PDF file
def extract_text_from_pdf(file_path):
    with open(file_path, 'rb') as file:
        reader = PyPDF2.PdfReader(file)
        text = ''
        for page in reader.pages:
            text += page.extract_text()
    return text

# Function to extract text from Word document
def extract_text_from_docx(file_path):
    doc = Document(file_path)
    text = ''
    for paragraph in doc.paragraphs:
        text += paragraph.text + '\n'
    return text

@app.route('/', methods=['GET', 'POST'])
@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        if 'files[]' not in request.files:
            return render_template('index.html', error="No files were uploaded.")
        
        files = request.files.getlist('files[]')
        cv_data = []

        for file in files:
            if file.filename == '':
                continue
            
            file_path = os.path.join(os.getcwd(), file.filename)
            file.save(file_path)
            
            file_name = file.filename
            file_ext = os.path.splitext(file_name)[1].lower()

            if file_ext == '.pdf':
                text = extract_text_from_pdf(file_path)
            elif file_ext == '.docx':
                text = extract_text_from_docx(file_path)
            else:
                continue
            
            emails = extract_emails(text)
            phone_numbers = extract_phone_numbers(text)
            
            cv_data.append((file_name, emails, phone_numbers, text))
        
            os.remove(file_path)  # Remove the temporary file
        
        # Create Excel workbook and write data
        wb = Workbook()
        ws = wb.active
        ws.append(["File Name", "Email", "Phone Number", "Text"])
        
        for data in cv_data:
            file_name, emails, phone_numbers, text = data
            for email, phone in zip(emails, phone_numbers):
                ws.append([file_name, email, phone, text])
        
        excel_file_path = 'cv_data.xlsx'
        wb.save(excel_file_path)
        
        return send_file(excel_file_path, as_attachment=True)
    
    return render_template('index.html')
if __name__ == '__main__':
    app.run(debug=True)
